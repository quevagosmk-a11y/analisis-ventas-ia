from __future__ import annotations

import io
from typing import Any, Callable, Dict, Tuple

from fpdf import FPDF
from openpyxl import Workbook


def normalize_export_format(raw_value: str | None) -> str:
    value = (raw_value or "pdf").strip().lower()
    if value == "excel":
        return "excel"
    if value == "pdf":
        return "pdf"
    raise ValueError("Formato de exportación inválido. Usa pdf o excel.")


def build_sales_report_export_file(
    payload: Dict[str, Any],
    export_format: str,
    *,
    pdf_safe_text: Callable[[Any], str],
    format_currency: Callable[[Any], str],
) -> Tuple[io.BytesIO, str, str]:
    period = payload.get("periodo") or {}
    ventas = payload.get("ventas") or []
    resumen = payload.get("resumen") or {}
    if export_format == "excel":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Ventas"
        sheet.append(["Reporte de ventas"])
        sheet.append(["Desde", period.get("desde") or "-"])
        sheet.append(["Hasta", period.get("hasta") or "-"])
        sheet.append(["Ventas registradas", resumen.get("ventas_registradas") or 0])
        sheet.append(["Productos vendidos", resumen.get("total_productos") or 0])
        sheet.append(["Monto total", float(resumen.get("monto_total") or 0.0)])
        sheet.append([])
        sheet.append(["ID", "Fecha", "Vendedor", "Metodo de pago", "Items", "Total"])
        for venta in ventas:
            sheet.append(
                [
                    int(venta.get("id") or 0),
                    venta.get("fecha") or "",
                    venta.get("vendedor") or "",
                    venta.get("metodo_pago") or "",
                    int(venta.get("items") or 0),
                    float(venta.get("total") or 0.0),
                ]
            )
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return (
            output,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    line_width = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, pdf_safe_text("Reporte de ventas"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(
        line_width,
        6,
        pdf_safe_text(
            f"Periodo: {period.get('desde') or '-'} a {period.get('hasta') or '-'}"
        ),
    )
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(
        line_width,
        6,
        pdf_safe_text(
            "Resumen: "
            f"{resumen.get('ventas_registradas') or 0} ventas, "
            f"{resumen.get('total_productos') or 0} productos, "
            f"total {format_currency(resumen.get('monto_total'))}"
        ),
    )
    if payload.get("offline"):
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(line_width, 6, pdf_safe_text("Fuente: respaldo offline"))
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, pdf_safe_text("Detalle"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    if not ventas:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(
            line_width,
            6,
            pdf_safe_text("No hay ventas para el periodo seleccionado."),
        )
    for venta in ventas:
        line = (
            f"#{int(venta.get('id') or 0)} | {venta.get('fecha') or '-'} | "
            f"{venta.get('vendedor') or '-'} | {venta.get('metodo_pago') or '-'} | "
            f"Items: {int(venta.get('items') or 0)} | "
            f"Total: {format_currency(venta.get('total'))}"
        )
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(line_width, 6, pdf_safe_text(line))
    output = io.BytesIO(bytes(pdf.output()))
    output.seek(0)
    return output, "application/pdf", "pdf"


def build_product_stats_export_file(
    payload: Dict[str, Any],
    export_format: str,
    *,
    pdf_safe_text: Callable[[Any], str],
    format_currency: Callable[[Any], str],
) -> Tuple[io.BytesIO, str, str]:
    products = payload.get("productos") or []
    summary = payload.get("resumen") or {}
    period_payload = payload.get("period") or {}
    period_label = payload.get("periodo") or "-"
    if export_format == "excel":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Estadisticas"
        sheet.append(["Reporte de estadisticas"])
        if period_payload:
            sheet.append(["Desde", period_payload.get("desde") or "-"])
            sheet.append(["Hasta", period_payload.get("hasta") or "-"])
        else:
            sheet.append(["Periodo", period_label])
        sheet.append(["Productos", summary.get("total_productos") or 0])
        sheet.append(["Unidades", summary.get("total_unidades") or 0])
        sheet.append(["Ingresos", float(summary.get("monto_total") or 0.0)])
        sheet.append([])
        sheet.append(["Producto", "Cantidad", "Ingresos"])
        for item in products:
            sheet.append(
                [
                    item.get("nombre") or "",
                    int(round(float(item.get("cantidad") or 0.0))),
                    float(item.get("ingresos") or 0.0),
                ]
            )
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return (
            output,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    line_width = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(
        0,
        10,
        pdf_safe_text("Reporte de estadisticas"),
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.set_font("Helvetica", size=10)
    if period_payload:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(
            line_width,
            6,
            pdf_safe_text(
                f"Periodo: {period_payload.get('desde') or '-'} a {period_payload.get('hasta') or '-'}"
            ),
        )
    else:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(line_width, 6, pdf_safe_text(f"Periodo: {period_label}"))
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(
        line_width,
        6,
        pdf_safe_text(
            "Resumen: "
            f"{summary.get('total_productos') or 0} productos, "
            f"{summary.get('total_unidades') or 0} unidades, "
            f"ingresos {format_currency(summary.get('monto_total'))}"
        ),
    )
    if payload.get("offline"):
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(line_width, 6, pdf_safe_text("Fuente: respaldo offline"))
    pdf.ln(2)
    pdf.set_font("Helvetica", size=9)
    if not products:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(
            line_width,
            6,
            pdf_safe_text("No hay estadisticas disponibles para el periodo seleccionado."),
        )
    for item in products:
        line = (
            f"{item.get('nombre') or '-'} | "
            f"Cantidad: {int(round(float(item.get('cantidad') or 0.0)))} | "
            f"Ingresos: {format_currency(item.get('ingresos'))}"
        )
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(line_width, 6, pdf_safe_text(line))
    output = io.BytesIO(bytes(pdf.output()))
    output.seek(0)
    return output, "application/pdf", "pdf"


def build_audit_export_file(
    payload: Dict[str, Any],
    export_format: str,
    *,
    pdf_safe_text: Callable[[Any], str],
    build_entity_label: Callable[[Dict[str, Any]], str],
    build_actor_label: Callable[[Dict[str, Any]], str],
) -> Tuple[io.BytesIO, str, str]:
    events = payload.get("eventos") or []
    summary = payload.get("summary") or {}
    period = summary.get("periodo") or {}
    if export_format == "excel":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Auditoria"
        sheet.append(["Reporte de auditoria"])
        if period:
            sheet.append(["Desde", period.get("desde") or "-"])
            sheet.append(["Hasta", period.get("hasta") or "-"])
        sheet.append(["Eventos", int(summary.get("total") or len(events) or 0)])
        if summary.get("action"):
            sheet.append(["Evento filtrado", summary.get("action")])
        if summary.get("user"):
            sheet.append(["Usuario filtrado", summary.get("user")])
        sheet.append([])
        sheet.append(["ID", "Fecha", "Evento", "Detalle", "Entidad", "Usuario"])
        for entry in events:
            sheet.append(
                [
                    int(entry.get("id") or 0),
                    entry.get("created_at") or "",
                    entry.get("event_type") or "",
                    entry.get("description") or "",
                    build_entity_label(entry),
                    build_actor_label(entry),
                ]
            )
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return (
            output,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    line_width = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(
        0, 10, pdf_safe_text("Reporte de auditoria"), new_x="LMARGIN", new_y="NEXT"
    )
    pdf.set_font("Helvetica", size=10)
    if period:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(
            line_width,
            6,
            pdf_safe_text(
                f"Periodo: {period.get('desde') or '-'} a {period.get('hasta') or '-'}"
            ),
        )
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(
        line_width,
        6,
        pdf_safe_text(f"Eventos: {int(summary.get('total') or len(events) or 0)}"),
    )
    if summary.get("action"):
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(
            line_width, 6, pdf_safe_text(f"Filtro evento: {summary.get('action')}")
        )
    if summary.get("user"):
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(
            line_width, 6, pdf_safe_text(f"Filtro usuario: {summary.get('user')}")
        )
    pdf.ln(2)
    pdf.set_font("Helvetica", size=9)
    if not events:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(
            line_width,
            6,
            pdf_safe_text("No hay eventos para el filtro seleccionado."),
        )
    for entry in events:
        line = (
            f"#{int(entry.get('id') or 0)} | {entry.get('created_at') or '-'} | "
            f"{entry.get('event_type') or '-'} | {build_entity_label(entry)} | "
            f"{build_actor_label(entry)} | {entry.get('description') or '-'}"
        )
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(line_width, 6, pdf_safe_text(line))
    output = io.BytesIO(bytes(pdf.output()))
    output.seek(0)
    return output, "application/pdf", "pdf"
