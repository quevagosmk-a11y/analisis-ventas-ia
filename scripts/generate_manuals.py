from __future__ import annotations

import json
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image, ImageDraw, ImageFont

Image.MAX_IMAGE_PIXELS = None


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs" / "manuales"
ASSETS_DIR = DOCS_DIR / "assets"
SCREENSHOTS_DIR = ASSETS_DIR / "screenshots"
DIAGRAMS_DIR = ASSETS_DIR / "diagrams"
OPTIMIZED_DIR = ASSETS_DIR / "optimized"
REFERENCE_DIR = DOCS_DIR / "referencia"
DOWNLOADS_DIR = Path.home() / "Downloads"

TECH_DOC_PATH = DOCS_DIR / "Manual_Tecnico_La_Septima_Estrella.docx"
USER_DOC_PATH = DOCS_DIR / "Manual_Usuario_La_Septima_Estrella.docx"
TECH_DOWNLOAD_PATH = DOWNLOADS_DIR / TECH_DOC_PATH.name
USER_DOWNLOAD_PATH = DOWNLOADS_DIR / USER_DOC_PATH.name

SYSTEM_NAME = "La Séptima Estrella"
PROJECT_NAME = "Sistema web de gestión operativa para inventario, ventas y analítica"
YEAR = "2026"


ROLE_ACCESS = {
    "Dashboard": "Administrador, gerente y vendedor.",
    "Registrar venta": "Administrador, gerente y vendedor con permiso de ventas.",
    "Productos": "Todos los perfiles con consulta; solo administrador y gerente con gestión completa.",
    "Reportes": "Administrador, gerente y vendedor.",
    "Estadísticas": "Administrador y gerente.",
    "Auditoría": "Administrador y auditador.",
    "Usuarios": "Administrador y gerente.",
    "Respaldos": "Administrador y gerente.",
    "Mi perfil": "Todos los perfiles autenticados.",
}


ENVIRONMENT_ROWS = [
    ("DB_HOST", "Host de MySQL", "127.0.0.1 en instalación local o mysql en Docker."),
    ("DB_USER", "Usuario de base de datos", "Cuenta con permisos de lectura y escritura sobre la base."),
    ("DB_PASSWORD", "Contraseña de base de datos", "Debe cambiarse antes de entrega o despliegue."),
    ("DB_NAME", "Nombre del esquema", "Por defecto la_septima_estrella."),
    ("APP_SECRET_KEY", "Clave de firma de sesión", "Debe ser larga, única y mantenerse fuera del código."),
    ("FLASK_DEBUG", "Modo depuración", "En entrega se mantiene en 0."),
    ("SESSION_IDLE_TIMEOUT_MINUTES", "Tiempo máximo de inactividad", "Controla la expiración por inactividad."),
    ("SESSION_MAX_AGE_DAYS", "Edad máxima de la sesión", "Límite de vigencia continua de la sesión."),
    ("DEFAULT_ADMIN_PASSWORD", "Clave semilla del administrador", "Solo se usa en bootstrap; debe cambiarse."),
    ("DEFAULT_VENDOR_PASSWORD", "Clave semilla del vendedor", "Solo se usa en bootstrap; debe cambiarse."),
    ("AI_ENGINE", "Motor del asistente", "Puede usar OpenAI, Ollama o heurística local."),
    ("AI_DEMAND_MODEL_PATH", "Ruta del modelo IA", "Apunta al modelo entrenado de pronóstico de demanda."),
]


PROJECT_STRUCTURE_ROWS = [
    ("src/main.py", "Punto de entrada Flask, registro de rutas, bootstrap y migraciones."),
    ("src/frontend/", "Capa de presentación: HTML, CSS y JavaScript de la interfaz web."),
    ("src/analysis/", "Lógica analítica de reportes, auditoría, dashboard y detalle de ventas."),
    ("src/ai/", "Motor del asistente, generación de features y pronóstico de demanda."),
    ("src/db/", "Gestión de roles, permisos y utilidades de esquema."),
    ("src/utils/", "Autenticación, sesiones, seguridad, respaldos y utilidades transversales."),
    ("src/data/", "Carga de ventas, datos demo y transformaciones de apoyo."),
    ("scripts/", "Automatizaciones de instalación, calidad, smoke tests y entrenamiento."),
    ("tests/", "Pruebas automatizadas funcionales, de permisos, sesiones y analítica."),
    ("docs/manuales/", "Salidas documentales, referencias y recursos visuales."),
]


API_ROWS = [
    ("POST", "/api/login", "Autenticar usuario y abrir sesión controlada.", "Público"),
    ("POST", "/api/logout", "Cerrar sesión activa.", "Usuario autenticado"),
    ("GET", "/api/dashboard", "Consultar indicadores operativos.", "Permiso dashboard_view"),
    ("GET", "/api/productos", "Consultar catálogo e inventario.", "Permiso products_view"),
    ("POST", "/api/ventas", "Registrar venta y descontar stock.", "Permiso sales_create"),
    ("GET", "/api/ventas/<id>", "Consultar detalle de una venta.", "Permiso sales_view"),
    ("POST", "/api/ventas/<id>/autorizar-deshabilitacion", "Habilitar deshabilitación controlada.", "Gerente o admin"),
    ("POST", "/api/ventas/<id>/deshabilitar", "Deshabilitar venta autorizada o directa.", "Vendedor autorizado o admin"),
    ("GET", "/api/reportes/ventas", "Consultar reporte de ventas por periodo.", "Permiso reports_view"),
    ("GET", "/api/reportes/ventas/export", "Exportar reporte en PDF o Excel.", "Permiso reports_view"),
    ("GET", "/api/estadisticas/productos-mas-vendidos", "Obtener estadísticas del periodo.", "Permiso statistics_view"),
    ("GET", "/api/auditoria", "Consultar eventos de trazabilidad.", "Permiso audit_view"),
    ("GET", "/api/backups", "Listar respaldos disponibles.", "Permiso backups_manage"),
    ("POST", "/api/backups/manual", "Generar respaldo manual del sistema.", "Permiso backups_manage"),
    ("POST", "/api/backups/restore", "Restaurar snapshot controlado.", "Permiso backups_manage"),
    ("GET", "/api/ia/pronostico-demanda", "Calcular demanda, cobertura y compra sugerida.", "Permiso statistics_view"),
    ("POST", "/api/ia/chat", "Consultar al asistente operativo.", "Permiso ai_chat"),
]


DATA_DICTIONARY_ROWS = [
    ("usuarios", "Credenciales, rol y estado operativo de cada cuenta."),
    ("roles_permisos", "Matriz de permisos por rol parametrizable."),
    ("sesiones_activas", "Control de tokens, actividad reciente y revocación."),
    ("productos", "Catálogo, precio, IVA, stock, mínimo y soporte de imagen."),
    ("inventario_movimientos", "Trazabilidad de entradas, salidas y ajustes de stock."),
    ("ventas", "Encabezado de venta, usuario, pago, total y estado de deshabilitación."),
    ("venta_detalle", "Líneas de producto por venta con cantidad, precio e IVA."),
    ("venta_recomendaciones", "Sugerencias de restock y combos ligadas a una venta."),
    ("producto_precio_auditoria", "Historial de cambios de precio y responsable."),
    ("auditoria_eventos", "Registro cronológico de acciones críticas del sistema."),
    ("schema_migrations", "Control de migraciones de esquema aplicadas."),
    ("configuracion_app / sistema_configuracion", "Parámetros operativos persistentes."),
]


QUALITY_ROWS = [
    ("Pruebas automatizadas", "106 pruebas", "Cobertura de autenticación, ventas, permisos, reportes, IA y respaldos."),
    ("Lint Python", "ruff limpio", "Sin hallazgos pendientes en la revisión estática."),
    ("Chequeo frontend", "Frontend JS OK", "Validación sintáctica del código JavaScript."),
    ("Smoke tests", "Playwright", "Recorrido automatizado sobre login y módulos principales."),
]


COMMON_MESSAGES = [
    ("Credenciales invalidas.", "Usuario o contraseña incorrectos.", "Verificar los datos e intentar nuevamente."),
    ("La sesión expiró.", "Se superó el tiempo de inactividad permitido.", "Iniciar sesión de nuevo."),
    ("Datos incompletos.", "Faltan campos obligatorios en el formulario.", "Completar los campos marcados y reenviar."),
    ("No se pudo registrar la venta.", "Error de validación, inventario o conexión.", "Revisar productos, cantidades y conexión antes de repetir."),
    ("No se pudo generar el reporte.", "El rango de fechas es inválido o hubo un error en la consulta.", "Validar el rango y volver a ejecutar."),
    ("No se pudo cargar el pronóstico.", "No hubo datos suficientes o falló la consulta analítica.", "Ajustar el periodo y reintentar."),
    ("No autorizado.", "El perfil activo no tiene permiso para la operación.", "Solicitar acceso al administrador o usar un perfil habilitado."),
    ("No se pudo restaurar el respaldo.", "Archivo inválido, inconsistencia o problema de base de datos.", "Intentar con otro snapshot o escalar a soporte."),
]


@dataclass(frozen=True)
class ScreenSection:
    title: str
    image_name: str
    purpose: str
    access: str
    fields: Sequence[tuple[str, str]]
    actions: Sequence[tuple[str, str]]
    flow: Sequence[str]
    validations: Sequence[str]


def ensure_dirs() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    DIAGRAMS_DIR.mkdir(parents=True, exist_ok=True)
    OPTIMIZED_DIR.mkdir(parents=True, exist_ok=True)


def get_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    preferred = "arialbd.ttf" if bold else "arial.ttf"
    windows_font = Path("C:/Windows/Fonts") / preferred
    if windows_font.exists():
        return ImageFont.truetype(str(windows_font), size=size)
    return ImageFont.load_default()


def set_cell_text(cell, text: str) -> None:
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10.5)


def format_table(table) -> None:
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10.5)


def set_default_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Arial"

    document.styles["Title"].font.size = Pt(16)
    document.styles["Subtitle"].font.size = Pt(12)
    document.styles["Heading 1"].font.size = Pt(14)
    document.styles["Heading 2"].font.size = Pt(13)
    document.styles["Heading 3"].font.size = Pt(12)

    for section in document.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(4)
        section.right_margin = Cm(2)
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)


def add_field(paragraph, instruction: str, default_text: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    if default_text:
        text = OxmlElement("w:t")
        text.text = default_text
        run._r.append(text)
    run._r.append(end)


def set_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Página ")
    add_field(paragraph, "PAGE", "1")
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)


def restart_page_numbering(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def add_cover_page(document: Document, title: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("[Nombre de la institución]").bold = True
    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("[Facultad y programa académico]")
    document.add_paragraph()
    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(PROJECT_NAME)
    document.add_paragraph()
    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Autor(es): [Nombre del autor o autores]")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Docente / director: [Nombre del docente o director]")
    document.add_paragraph()
    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("[Ciudad]")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(YEAR)


def start_main_section(document: Document):
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.different_first_page_header_footer = False
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)
    section.left_margin = Cm(4)
    section.right_margin = Cm(2)
    restart_page_numbering(section, 1)
    set_page_number(section)
    return section


def add_toc(document: Document) -> None:
    document.add_heading("Tabla de contenido", level=1)
    p = document.add_paragraph()
    add_field(p, r'TOC \o "1-3" \h \z \u', "Actualice la tabla en Word.")
    note = document.add_paragraph(
        "Nota: al abrir el documento en Microsoft Word, debe actualizarse la tabla de contenido para reflejar la numeración definitiva."
    )
    note.italic = True


def add_paragraph(document: Document, text: str, *, align: int | None = None) -> None:
    p = document.add_paragraph(text)
    if align is not None:
        p.alignment = align


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: Iterable[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_table(document: Document, headers: Sequence[str], rows: Sequence[Sequence[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    format_table(table)


def prepare_image(path: Path, max_width_px: int = 1600, max_height_px: int = 2500) -> Path:
    OPTIMIZED_DIR.mkdir(parents=True, exist_ok=True)
    target = OPTIMIZED_DIR / path.name
    with Image.open(path) as image:
        image = image.convert("RGB")
        resize_ratio = min(
            max_width_px / float(image.width),
            max_height_px / float(image.height),
            1.0,
        )
        if resize_ratio < 1.0:
            image = image.resize(
                (
                    max(1, int(image.width * resize_ratio)),
                    max(1, int(image.height * resize_ratio)),
                ),
                Image.LANCZOS,
            )
        image.save(target, format="JPEG", quality=88, optimize=True)
    return target


def add_figure(document: Document, image_path: Path, caption: str, number: int) -> None:
    prepared = prepare_image(image_path)
    with Image.open(prepared) as image:
        max_width = 6.0 if image.width >= image.height else 4.8
    document.add_picture(str(prepared), width=Inches(max_width))
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Figura {number}. {caption}")
    run.italic = True


def write_centered_text(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, font, fill: str) -> None:
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=4)
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(
        (xy[0] - width / 2, xy[1] - height / 2),
        text,
        font=font,
        fill=fill,
        spacing=4,
        align="center",
    )


def draw_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, body: str, fill: str) -> None:
    draw.rounded_rectangle(box, radius=18, fill=fill, outline="#233b83", width=3)
    title_font = get_font(26, bold=True)
    body_font = get_font(18)
    x1, y1, x2, y2 = box
    write_centered_text(draw, ((x1 + x2) // 2, y1 + 40), title, title_font, "#0d1b4d")
    write_centered_text(draw, ((x1 + x2) // 2, (y1 + y2) // 2 + 10), body, body_font, "#132042")


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str = "#233b83") -> None:
    draw.line((start, end), fill=fill, width=6)
    arrow_size = 14
    if end[0] >= start[0]:
        points = [end, (end[0] - arrow_size, end[1] - arrow_size // 2), (end[0] - arrow_size, end[1] + arrow_size // 2)]
    else:
        points = [end, (end[0] + arrow_size, end[1] - arrow_size // 2), (end[0] + arrow_size, end[1] + arrow_size // 2)]
    draw.polygon(points, fill=fill)


def create_architecture_diagram() -> Path:
    path = DIAGRAMS_DIR / "arquitectura_general.png"
    if path.exists():
        return path
    image = Image.new("RGB", (1600, 900), "#f5f7ff")
    draw = ImageDraw.Draw(image)
    header_font = get_font(34, bold=True)
    write_centered_text(draw, (800, 70), "Arquitectura general del sistema", header_font, "#0d1b4d")
    draw_box(draw, (90, 180, 470, 450), "Capa de presentación", "HTML + CSS + JavaScript\nPanel web responsivo\nNavegación por módulos", "#dfe9ff")
    draw_box(draw, (610, 140, 1010, 490), "Backend Flask", "Rutas REST internas\nAutenticación y sesiones\nVentas, reportes, IA y respaldos", "#d8f0ff")
    draw_box(draw, (1140, 120, 1490, 330), "Base de datos", "MySQL relacional\nIntegridad y auditoría", "#e7f6e8")
    draw_box(draw, (1140, 390, 1490, 600), "Persistencia auxiliar", "Snapshots JSON\nReportes PDF / Excel", "#fff2d8")
    draw_box(draw, (610, 560, 1010, 790), "Servicios analíticos", "Pronóstico de demanda\nReglas heurísticas\nModelo entrenado", "#ffe3ea")
    draw_box(draw, (90, 560, 470, 790), "Usuarios del negocio", "Administrador\nGerente\nVendedor\nAuditador", "#efe2ff")
    draw_arrow(draw, (470, 315), (610, 315))
    draw_arrow(draw, (1010, 250), (1140, 225))
    draw_arrow(draw, (1010, 430), (1140, 495))
    draw_arrow(draw, (800, 490), (800, 560))
    draw_arrow(draw, (470, 675), (610, 675))
    image.save(path)
    return path


def create_navigation_diagram() -> Path:
    path = DIAGRAMS_DIR / "mapa_navegacion.png"
    if path.exists():
        return path
    image = Image.new("RGB", (1600, 1000), "#fbfcff")
    draw = ImageDraw.Draw(image)
    header_font = get_font(34, bold=True)
    write_centered_text(draw, (800, 60), "Mapa de navegación funcional", header_font, "#0d1b4d")
    center_box = (620, 110, 980, 220)
    draw_box(draw, center_box, "Inicio de sesión", "Acceso, validación de credenciales\ny apertura de sesión", "#e1eaff")
    modules = [
        ((110, 330, 450, 460), "Dashboard", "Indicadores y resumen"),
        ((510, 330, 850, 460), "Registrar venta", "Búsqueda, carrito y cierre"),
        ((910, 330, 1250, 460), "Productos", "Consulta e inventario"),
        ((1310, 330, 1540, 460), "Reportes", "Consulta y exportación"),
        ((110, 620, 450, 750), "Estadísticas", "Tendencias y pronóstico IA"),
        ((510, 620, 850, 750), "Usuarios", "Roles y permisos"),
        ((910, 620, 1250, 750), "Auditoría", "Trazabilidad y filtros"),
        ((1310, 620, 1540, 750), "Respaldos / Perfil", "Backups, restauración y datos de cuenta"),
    ]
    for box, title, body in modules:
        draw_box(draw, box, title, body, "#eef3ff")
        draw_arrow(draw, ((center_box[0] + center_box[2]) // 2, center_box[3]), ((box[0] + box[2]) // 2, box[1]))
    image.save(path)
    return path


def create_data_model_diagram() -> Path:
    path = DIAGRAMS_DIR / "modelo_datos_resumen.png"
    if path.exists():
        return path
    image = Image.new("RGB", (1600, 1000), "#f8fbff")
    draw = ImageDraw.Draw(image)
    header_font = get_font(34, bold=True)
    write_centered_text(draw, (800, 60), "Resumen del modelo de datos", header_font, "#0d1b4d")
    boxes = {
        "usuarios": (120, 160, 460, 340),
        "sesiones_activas": (120, 420, 460, 600),
        "auditoria_eventos": (120, 690, 460, 870),
        "ventas": (620, 140, 980, 330),
        "venta_detalle": (1080, 180, 1460, 360),
        "productos": (1080, 520, 1460, 760),
        "inventario_movimientos": (620, 500, 980, 720),
    }
    labels = {
        "usuarios": ("usuarios", "credenciales, rol\ny estado"),
        "sesiones_activas": ("sesiones_activas", "token, actividad,\nrevocación"),
        "auditoria_eventos": ("auditoria_eventos", "evento, actor,\nfecha"),
        "ventas": ("ventas", "usuario, fecha,\ntotal y estado"),
        "venta_detalle": ("venta_detalle", "productos vendidos,\ncantidad e IVA"),
        "productos": ("productos", "catálogo, stock,\nprecio e imagen"),
        "inventario_movimientos": ("inventario_movimientos", "entradas, salidas\ny ajustes"),
    }
    for key, box in boxes.items():
        draw_box(draw, box, labels[key][0], labels[key][1], "#e8f0ff")
    draw_arrow(draw, (460, 250), (620, 235))
    draw_arrow(draw, (980, 235), (1080, 270))
    draw_arrow(draw, (1270, 360), (1270, 520))
    draw_arrow(draw, (1080, 640), (980, 640))
    draw_arrow(draw, (290, 340), (290, 420))
    draw_arrow(draw, (290, 600), (290, 690))
    draw_arrow(draw, (460, 250), (620, 620))
    image.save(path)
    return path


def create_voiding_flow_diagram() -> Path:
    path = DIAGRAMS_DIR / "flujo_deshabilitacion_venta.png"
    if path.exists():
        return path
    image = Image.new("RGB", (1600, 850), "#fffdf8")
    draw = ImageDraw.Draw(image)
    header_font = get_font(34, bold=True)
    write_centered_text(draw, (800, 60), "Flujo de deshabilitación controlada de ventas", header_font, "#6b3d00")
    steps = [
        ((80, 260, 360, 480), "1. Venta registrada", "La venta queda activa\ny afecta inventario."),
        ((430, 260, 720, 480), "2. Habilitación puntual", "El gerente o admin\nvalida la clave borrar."),
        ((790, 260, 1080, 480), "3. Deshabilitación", "El vendedor autorizado\no el admin ejecuta la acción."),
        ((1150, 260, 1520, 480), "4. Trazabilidad", "La venta no se elimina;\nse revierte stock y queda auditada."),
    ]
    for box, title, body in steps:
        draw_box(draw, box, title, body, "#fff0cf")
    draw_arrow(draw, (360, 370), (430, 370), fill="#6b3d00")
    draw_arrow(draw, (720, 370), (790, 370), fill="#6b3d00")
    draw_arrow(draw, (1080, 370), (1150, 370), fill="#6b3d00")
    image.save(path)
    return path


def generate_diagrams() -> dict[str, Path]:
    return {
        "architecture": create_architecture_diagram(),
        "navigation": create_navigation_diagram(),
        "data_model": create_data_model_diagram(),
        "sale_voiding": create_voiding_flow_diagram(),
    }


SCREEN_SECTIONS = [
    ScreenSection(
        title="Pantalla de inicio de sesión",
        image_name="01_login.png",
        purpose="Permite autenticar al usuario y abrir una sesión controlada según el rol configurado.",
        access="Pantalla pública previa a la autenticación.",
        fields=[("Usuario", "Identificador de la cuenta registrada."), ("Contraseña", "Clave asociada a la cuenta.")],
        actions=[("Ingresar", "Valida credenciales, permisos y estado de la cuenta.")],
        flow=[
            "El usuario digita su nombre de usuario y contraseña.",
            "El sistema valida los datos y crea una sesión activa.",
            "La aplicación redirige al dashboard del perfil autenticado.",
        ],
        validations=[
            "Si faltan datos, el sistema informa que existen campos incompletos.",
            "Si hay múltiples intentos fallidos, se aplica bloqueo temporal.",
        ],
    ),
    ScreenSection(
        title="Dashboard operativo",
        image_name="02_dashboard.png",
        purpose="Resume ventas, alertas de inventario, actividad reciente y accesos directos por módulo.",
        access=ROLE_ACCESS["Dashboard"],
        fields=[
            ("Tarjetas de indicador", "Muestran totales y variaciones del periodo."),
            ("Paneles de actividad", "Presentan ventas recientes, alertas y resumen del sistema."),
        ],
        actions=[
            ("Navegación lateral", "Permite desplazarse a los módulos habilitados por el rol."),
        ],
        flow=[
            "El sistema carga indicadores del día o del periodo configurado.",
            "El usuario revisa el estado general y accede al módulo que necesita operar.",
        ],
        validations=[
            "Si no existe conexión a la base, la interfaz informa el uso de respaldo local.",
        ],
    ),
    ScreenSection(
        title="Registro de ventas",
        image_name="03_sales.png",
        purpose="Permite buscar productos, armar un carrito, seleccionar método de pago y confirmar la transacción.",
        access=ROLE_ACCESS["Registrar venta"],
        fields=[
            ("Buscador de productos", "Filtra el catálogo disponible por nombre o categoría."),
            ("Tabla de resultados", "Muestra precio, stock y acción para agregar."),
            ("Carrito", "Consolida productos, cantidades, IVA y total."),
            ("Método de pago", "Selecciona la forma de recaudo de la venta."),
        ],
        actions=[
            ("Agregar", "Envía el producto al carrito."),
            ("Registrar venta", "Confirma la operación y descuenta stock."),
            ("Limpiar carrito", "Vacía la selección actual."),
        ],
        flow=[
            "El usuario busca uno o más productos.",
            "Agrega las unidades requeridas al carrito.",
            "Selecciona el método de pago y registra la venta.",
            "El sistema confirma, actualiza inventario y registra auditoría cuando aplica.",
        ],
        validations=[
            "No permite registrar cantidades mayores al stock disponible.",
            "Rechaza carritos vacíos o cantidades inválidas.",
        ],
    ),
    ScreenSection(
        title="Detalle de venta y deshabilitación controlada",
        image_name="04_sale_detail.png",
        purpose="Muestra el detalle completo de una venta y permite, según el rol, habilitar o deshabilitar una operación errónea.",
        access="Consulta disponible para perfiles con ventas. La habilitación corresponde a gerente o admin; la deshabilitación al vendedor autorizado o admin.",
        fields=[
            ("Encabezado de venta", "Identificador, fecha, usuario, método de pago y total."),
            ("Detalle por producto", "Lista de productos, cantidades y precios involucrados."),
            ("Estado", "Indica si la venta está activa, habilitada o deshabilitada."),
        ],
        actions=[
            ("Habilitar deshabilitación", "Autoriza puntualmente la venta con la clave borrar."),
            ("Deshabilitar venta", "Revierte la venta sin eliminar su trazabilidad."),
        ],
        flow=[
            "El gerente revisa la venta que presentó error.",
            "Si corresponde, habilita solo esa venta mediante la clave de autorización.",
            "El vendedor autorizado entra al detalle y ejecuta la deshabilitación.",
            "El sistema revierte stock, excluye la venta de totales y deja auditoría.",
        ],
        validations=[
            "No permite deshabilitar una venta ya deshabilitada.",
            "Un vendedor no puede deshabilitar una venta no autorizada.",
        ],
    ),
    ScreenSection(
        title="Gestión de productos",
        image_name="05_products.png",
        purpose="Administra el catálogo, niveles de stock, imágenes, IVA y estado activo de cada producto.",
        access=ROLE_ACCESS["Productos"],
        fields=[
            ("Formulario de producto", "Nombre, categoría, precio, IVA, stock, mínimo e imagen."),
            ("Tabla de productos", "Resume la información comercial y el inventario actual."),
            ("Filtro / búsqueda", "Localiza productos por nombre o categoría."),
        ],
        actions=[
            ("Guardar", "Crea o actualiza un producto."),
            ("Activar / desactivar", "Controla si el producto puede venderse."),
            ("Exportar / importar", "Maneja inventario masivo en CSV o Excel."),
        ],
        flow=[
            "El usuario con gestión completa diligencia o actualiza los datos del producto.",
            "El sistema guarda la información y refresca la tabla.",
        ],
        validations=[
            "Valida precio, stock, IVA y mínimo como datos numéricos.",
            "Controla formatos de imagen y evita datos obligatorios vacíos.",
        ],
    ),
    ScreenSection(
        title="Reportes de ventas",
        image_name="06_reports.png",
        purpose="Consulta ventas por rango de fechas y permite exportar los resultados en PDF o Excel.",
        access=ROLE_ACCESS["Reportes"],
        fields=[
            ("Fecha desde", "Inicio del periodo a consultar."),
            ("Fecha hasta", "Fin del periodo a consultar."),
            ("Resumen del periodo", "Cantidad de ventas, productos vendidos y total generado."),
            ("Tabla de resultados", "Detalle de ventas incluidas en el rango."),
        ],
        actions=[
            ("Generar", "Consulta el reporte en pantalla."),
            ("Exportar PDF", "Genera una salida imprimible del periodo."),
            ("Exportar Excel", "Entrega el reporte en formato editable."),
        ],
        flow=[
            "El usuario define el rango de fechas.",
            "El sistema muestra el resumen y la tabla del periodo.",
            "Si se requiere, el reporte se exporta en el formato seleccionado.",
        ],
        validations=[
            "No acepta rangos incompletos.",
            "La fecha inicial no puede ser mayor que la fecha final.",
        ],
    ),
    ScreenSection(
        title="Estadísticas y pronóstico de demanda",
        image_name="08_forecast.png",
        purpose="Entrega indicadores de productos más vendidos y estimaciones de demanda, cobertura y compra sugerida.",
        access=ROLE_ACCESS["Estadísticas"],
        fields=[
            ("Periodo analizado", "Define la base histórica de las estadísticas."),
            ("Tarjetas de resumen", "Muestran productos evaluados, riesgo crítico y cobertura promedio."),
            ("Tabla de pronóstico", "Indica pronóstico, cobertura, riesgo y compra sugerida por producto."),
        ],
        actions=[
            ("Generar estadísticas", "Procesa el periodo seleccionado."),
            ("Exportar PDF / Excel", "Descarga resultados para análisis externo."),
        ],
        flow=[
            "El usuario establece el periodo de observación.",
            "El sistema calcula indicadores históricos y pronóstico IA.",
            "La tabla resalta productos críticos y necesidades de reposición.",
        ],
        validations=[
            "Si no existen datos suficientes, el sistema usa heurística controlada.",
            "El cálculo del horizonte no muestra fechas de agotamiento anteriores al día actual.",
        ],
    ),
    ScreenSection(
        title="Gestión de usuarios y roles",
        image_name="09_users.png",
        purpose="Permite crear, editar, activar o desactivar usuarios y administrar permisos por rol.",
        access=ROLE_ACCESS["Usuarios"],
        fields=[
            ("Formulario de usuario", "Nombre, usuario, rol, estado y permisos."),
            ("Tabla de usuarios", "Lista las cuentas registradas y su estado."),
        ],
        actions=[
            ("Guardar usuario", "Crea o actualiza una cuenta."),
            ("Editar", "Carga la información en el formulario."),
            ("Activar / desactivar", "Controla la disponibilidad de acceso."),
        ],
        flow=[
            "El administrador o gerente registra una nueva cuenta.",
            "El sistema asigna rol y permisos, y la cuenta queda disponible para inicio de sesión.",
        ],
        validations=[
            "No permite roles vacíos ni usuarios duplicados.",
            "Requiere al menos un permiso funcional cuando se crean roles personalizados.",
        ],
    ),
    ScreenSection(
        title="Auditoría del sistema",
        image_name="10_audit.png",
        purpose="Consulta la trazabilidad de eventos relevantes como inicios de sesión, ventas, cambios de precio y respaldos.",
        access=ROLE_ACCESS["Auditoría"],
        fields=[
            ("Filtros de fecha y tipo", "Acotan el conjunto de eventos consultados."),
            ("Tabla de auditoría", "Muestra fecha, evento, detalle, entidad y actor."),
        ],
        actions=[
            ("Filtrar", "Recarga la auditoría según criterios."),
            ("Actualizar", "Restablece la consulta."),
            ("Exportar PDF / Excel", "Descarga la trazabilidad del periodo."),
        ],
        flow=[
            "El usuario autorizado define filtros.",
            "El sistema consulta eventos y los presenta en orden cronológico.",
        ],
        validations=[
            "Si la consulta falla, la pantalla informa que no se pudo cargar la auditoría.",
        ],
    ),
    ScreenSection(
        title="Respaldos del sistema",
        image_name="11_backups.png",
        purpose="Lista snapshots generados, permite descargar evidencias y restaurar un estado previo del sistema.",
        access=ROLE_ACCESS["Respaldos"],
        fields=[
            ("Lista de respaldos", "Nombre del archivo, fecha y rutas de descarga."),
            ("Mensaje operativo", "Informa si el respaldo se generó o restauró correctamente."),
        ],
        actions=[
            ("Generar respaldo", "Crea un snapshot manual."),
            ("Descargar JSON / PDF", "Descarga la evidencia del respaldo."),
            ("Restaurar", "Recupera la información almacenada en el snapshot."),
        ],
        flow=[
            "El usuario autorizado genera o selecciona un respaldo existente.",
            "Si restaura, el sistema reemplaza datos de productos, ventas e inventario con trazabilidad.",
        ],
        validations=[
            "Solo acepta rutas internas válidas de respaldo.",
            "Si el archivo es inválido, se informa el error y no se restaura nada.",
        ],
    ),
    ScreenSection(
        title="Mi perfil y sesión",
        image_name="12_account.png",
        purpose="Permite revisar datos básicos de la cuenta, capacidades disponibles y cerrar sesión de forma segura.",
        access=ROLE_ACCESS["Mi perfil"],
        fields=[
            ("Datos del perfil", "Nombre, usuario, rol y permisos activos."),
            ("Acciones de cuenta", "Opciones de cierre de sesión y, según rol, generación de respaldo."),
        ],
        actions=[
            ("Cerrar sesión", "Revoca la sesión activa."),
            ("Generar respaldo", "Disponible para perfiles con permiso sobre respaldos."),
        ],
        flow=[
            "El usuario consulta su perfil y capacidades.",
            "Cuando finaliza la jornada, cierra sesión para liberar el acceso.",
        ],
        validations=[
            "Si la sesión expira, el sistema redirige al login.",
        ],
    ),
]


def add_screen_section(document: Document, section: ScreenSection, figure_counter: list[int]) -> None:
    document.add_heading(section.title, level=2)
    add_paragraph(document, section.purpose)
    add_paragraph(document, f"Acceso: {section.access}")
    figure_counter[0] += 1
    add_figure(document, SCREENSHOTS_DIR / section.image_name, section.title, figure_counter[0])
    document.add_paragraph("Campos presentes en la pantalla:")
    add_table(document, ["Campo / zona", "Descripción"], section.fields)
    document.add_paragraph("Botones y acciones disponibles:")
    add_table(document, ["Elemento", "Descripción"], section.actions)
    document.add_paragraph("Flujo normal de uso:")
    add_numbered(document, section.flow)
    document.add_paragraph("Validaciones y controles:")
    add_bullets(document, section.validations)


def read_reference_excerpt(filename: str) -> str:
    path = REFERENCE_DIR / filename
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8", errors="ignore")


def build_technical_manual(diagrams: dict[str, Path]) -> None:
    doc = Document()
    set_default_styles(doc)
    add_cover_page(doc, f"MANUAL TÉCNICO DEL SISTEMA {SYSTEM_NAME.upper()}")
    start_main_section(doc)
    add_toc(doc)
    doc.add_page_break()

    figure_counter = [0]

    doc.add_heading("Presentación", level=1)
    add_paragraph(
        doc,
        "El presente manual técnico documenta la instalación, configuración, operación, mantenimiento y evolución del sistema web La Séptima Estrella. El documento se dirige a personal con formación técnica en desarrollo, soporte, administración de sistemas y bases de datos. Su propósito consiste en facilitar la comprensión integral del producto, su arquitectura y sus mecanismos de operación en un entorno académico o institucional.",
    )

    doc.add_heading("Resumen ejecutivo", level=1)
    add_paragraph(
        doc,
        "La Séptima Estrella es un sistema web de gestión operativa orientado al control de inventario, ventas, usuarios, reportes, auditoría, respaldos y analítica con apoyo de inteligencia artificial. La solución adopta una arquitectura web cliente-servidor con una interfaz en HTML, CSS y JavaScript, un backend monolítico modularizado sobre Flask y una base de datos relacional MySQL. El manual se estructura en seis capítulos principales: requisitos, diagramas, aspectos técnicos del desarrollo, instalación, requerimientos de operación y referencias. Cada capítulo permite al lector técnico comprender la solución desde la infraestructura mínima hasta su mantenimiento evolutivo.",
    )

    doc.add_heading("Finalidad del manual", level=1)
    add_bullets(
        doc,
        [
            "Capacitar al personal técnico sobre instalación y puesta en marcha.",
            "Servir como referencia de mantenimiento correctivo y evolutivo.",
            "Documentar la estructura lógica, los módulos y el modelo de datos.",
            "Estandarizar procedimientos de respaldo, restauración y contingencia.",
        ],
    )

    doc.add_heading("Introducción", level=1)
    add_paragraph(
        doc,
        "El documento se organiza siguiendo el instructivo oficial suministrado para manuales técnicos. En el primer capítulo se exponen requisitos de hardware, software y contingencia. El segundo capítulo presenta diagramas útiles para comprender la arquitectura, la navegación y el modelo de datos. El tercero describe la capa de presentación, la capa de servicios, la base de datos y los mecanismos de seguridad. El cuarto capítulo detalla la instalación del sistema y sus servicios complementarios. El quinto expone requerimientos operativos y recomendaciones de escalabilidad. Finalmente, el sexto capítulo reúne referencias técnicas consultadas para el desarrollo y la validación del sistema.",
    )

    doc.add_heading("1. Requisitos del sistema", level=1)
    doc.add_heading("1.1 Requisitos de hardware", level=2)
    add_table(
        doc,
        ["Nivel", "Procesador", "Memoria RAM", "Almacenamiento", "Observación"],
        [
            ("Mínimo", "Intel i3 / Ryzen 3 o equivalente", "8 GB", "2 GB libres para proyecto y snapshots", "Adecuado para entorno académico o de pruebas."),
            ("Recomendado", "Intel i5 / Ryzen 5 o superior", "16 GB", "10 GB libres", "Mejor respuesta para pruebas, entrenamiento y Docker."),
        ],
    )
    doc.add_heading("1.2 Requisitos de software", level=2)
    add_table(
        doc,
        ["Componente", "Versión o referencia", "Uso dentro del sistema"],
        [
            ("Sistema operativo", "Windows 10/11, Linux o WSL", "Ejecución local del proyecto."),
            ("Python", "3.12 o superior", "Backend Flask, scripts y pruebas."),
            ("Node.js", "Solo para chequeos del frontend", "Validaciones y smoke checks complementarios."),
            ("MySQL", "8.x", "Persistencia principal del sistema."),
            ("Docker Desktop", "Opcional", "Despliegue simplificado con app + MySQL."),
            ("Navegador", "Chrome, Edge o equivalente moderno", "Consumo de la interfaz web."),
        ],
    )
    add_paragraph(doc, "Dependencias destacadas del proyecto:")
    add_bullets(
        doc,
        [
            "Flask y Flask-Cors para el backend HTTP.",
            "PyMySQL para acceso a MySQL.",
            "pandas, numpy y scikit-learn para analítica y pronóstico.",
            "fpdf2 y openpyxl para exportación de reportes.",
            "python-dotenv para configuración por entorno.",
        ],
    )
    doc.add_heading("1.3 Configuraciones previas", level=2)
    add_table(doc, ["Variable", "Propósito", "Observación"], ENVIRONMENT_ROWS)
    doc.add_heading("1.4 Portabilidad y adaptabilidad", level=2)
    add_paragraph(
        doc,
        "El sistema puede ejecutarse en entorno local tradicional o en contenedores Docker. La aplicación se apoya en variables de entorno para desacoplar configuraciones sensibles, lo que facilita su adaptación entre equipos de desarrollo, entornos de evaluación y despliegues controlados. El proyecto también incluye scripts de bootstrap, instalación automatizada y sincronización del catálogo, lo que reduce la dependencia de configuraciones manuales repetitivas.",
    )
    doc.add_heading("1.5 Contingencia", level=2)
    add_bullets(
        doc,
        [
            "El sistema genera snapshots manuales y automáticos en la carpeta backups/.",
            "Cuando MySQL no está disponible, la aplicación puede cargar el último respaldo y operar en modo resiliente de solo lectura.",
            "El módulo de respaldos permite descargar evidencia en JSON y PDF, así como restaurar el estado persistido.",
            "Las migraciones de esquema se controlan mediante la tabla schema_migrations.",
        ],
    )

    doc.add_heading("2. Diagramas de modelamiento", level=1)
    add_paragraph(
        doc,
        "Conforme al instructivo, se incluyen diagramas que aportan valor para la comprensión del sistema. Cada representación se acompaña de una explicación breve sobre su propósito y relación con la solución desarrollada.",
    )
    figure_counter[0] += 1
    add_figure(doc, diagrams["architecture"], "Arquitectura general en capas del sistema", figure_counter[0])
    add_paragraph(doc, "La figura anterior resume la relación entre la interfaz web, el backend Flask, la persistencia relacional, el almacenamiento de respaldos y los servicios analíticos.")
    figure_counter[0] += 1
    add_figure(doc, diagrams["navigation"], "Mapa de navegación funcional del sistema", figure_counter[0])
    add_paragraph(doc, "El mapa de navegación facilita identificar la pantalla inicial, los módulos principales y la relación entre opciones del menú.")
    figure_counter[0] += 1
    add_figure(doc, diagrams["data_model"], "Resumen del modelo de datos relacional", figure_counter[0])
    add_paragraph(doc, "El diagrama resume las entidades más relevantes para ventas, inventario, sesiones y auditoría.")
    figure_counter[0] += 1
    add_figure(doc, diagrams["sale_voiding"], "Flujo controlado de deshabilitación de ventas", figure_counter[0])
    add_paragraph(doc, "Este flujo documenta una regla de negocio crítica: la venta no se elimina de la base, sino que se deshabilita con trazabilidad y restitución de stock.")

    doc.add_heading("3. Aspectos técnicos del desarrollo del sistema", level=1)
    doc.add_heading("3.1 Arquitectura del sistema", level=2)
    add_paragraph(
        doc,
        "La solución se implementa como una aplicación web con arquitectura cliente-servidor. Aunque el punto de entrada principal se mantiene en un backend monolítico, el código fue modularizado por responsabilidades para reducir acoplamiento y facilitar mantenimiento. La interfaz consume rutas internas tipo API y presenta la información en una experiencia tipo panel administrativo.",
    )
    doc.add_heading("a. Capa de presentación", level=3)
    add_paragraph(
        doc,
        "La capa de presentación utiliza HTML, CSS y JavaScript. Su objetivo consiste en mostrar información operativa, guiar al usuario por módulos definidos por permisos y consumir datos del backend mediante solicitudes HTTP. El frontend valida rangos de fechas, estados de formularios, exportaciones, errores operativos y sesiones expiradas.",
    )
    add_paragraph(doc, "Relación de directorios de la capa de presentación:")
    add_table(doc, ["Ruta", "Descripción"], [("src/frontend/index.html", "Estructura principal de la aplicación."), ("src/frontend/app.js", "Lógica de interacción, consumo API y renderizado."), ("src/frontend/style.css", "Estilos visuales y diseño responsivo."), ("src/frontend/img/", "Recursos gráficos de la interfaz.")])
    add_paragraph(doc, "Mapa de navegación o estructura del sistema:")
    add_bullets(doc, [f"{name}: {desc}" for name, desc in ROLE_ACCESS.items()])
    add_paragraph(doc, "Archivos de configuración relacionados con la presentación:")
    add_table(doc, ["Archivo", "Propósito"], [(".env / .env.example", "Define parámetros del entorno, seguridad y servicios."), ("package.json", "Declara scripts de validación del frontend."), ("pyproject.toml", "Configura herramientas de calidad del proyecto.")])

    doc.add_heading("b. Capa de lógica y servicios", level=3)
    add_paragraph(
        doc,
        "La capa de lógica se encuentra sobre Flask y coordina autenticación, autorización, sesiones, ventas, productos, reportes, auditoría, respaldos y analítica. Las rutas se apoyan en módulos de analysis, utils, ai y db para encapsular responsabilidades concretas. El backend controla transacciones al registrar ventas, aplica auditoría ante acciones sensibles y valida permisos por rol antes de resolver una operación.",
    )
    add_paragraph(doc, "Estructura técnica relevante del proyecto:")
    add_table(doc, ["Ruta", "Responsabilidad"], PROJECT_STRUCTURE_ROWS)
    add_paragraph(doc, "Principales endpoints internos del sistema:")
    add_table(doc, ["Método", "Ruta", "Finalidad", "Acceso"], API_ROWS)

    doc.add_heading("c. Capa de base de datos", level=3)
    add_paragraph(
        doc,
        "La persistencia principal se implementa en una base de datos relacional MySQL. El diseño utiliza claves primarias y foráneas para garantizar integridad referencial entre usuarios, ventas, detalle de ventas, productos, movimientos de inventario, auditoría y sesiones. Las migraciones incrementales permiten adaptar el esquema sin reconstruir la base manualmente.",
    )
    add_paragraph(doc, "Diccionario de datos resumido:")
    add_table(doc, ["Tabla", "Descripción funcional"], DATA_DICTIONARY_ROWS)

    doc.add_heading("3.2 Seguridad, validaciones y trazabilidad", level=2)
    add_bullets(
        doc,
        [
            "Las contraseñas se almacenan con PBKDF2-SHA256 y el sistema puede migrar hashes antiguos al iniciar sesión.",
            "La sesión expira por inactividad y por edad máxima configurada.",
            "Existe bloqueo temporal ante múltiples intentos fallidos de autenticación.",
            "Las rutas verifican permisos funcionales por rol antes de responder.",
            "Los cambios de precio, respaldos, aperturas y cierres de sesión, y la deshabilitación de ventas quedan registrados en auditoría.",
        ],
    )

    doc.add_heading("3.3 Inteligencia artificial y pronóstico de demanda", level=2)
    add_paragraph(
        doc,
        "El módulo de inteligencia artificial utiliza un enfoque híbrido. Cuando existe un modelo entrenado, el backend carga un RandomForestRegressor persistido en models/demand_model.joblib. Cuando no existe modelo o los datos son insuficientes, el sistema recurre a una heurística controlada basada en histórico y ventana temporal real. El resultado incluye pronóstico de demanda, promedio diario, cobertura, riesgo de quiebre y compra sugerida por producto. El frontend presenta estos indicadores en la sección de estadísticas, junto con métricas de calidad del modelo cuando están disponibles.",
    )

    doc.add_heading("3.4 Pruebas y control de calidad", level=2)
    add_paragraph(
        doc,
        "La verificación del sistema se apoya en pruebas automatizadas, revisión estática del código y smoke tests de interfaz. A la fecha de elaboración del manual, la base de validación reporta 106 pruebas exitosas, ruff sin hallazgos y chequeo de frontend satisfactorio.",
    )
    add_table(doc, ["Actividad", "Resultado", "Observación"], QUALITY_ROWS)

    doc.add_heading("4. Instalación del sistema y servicios", level=1)
    doc.add_heading("4.1 Instalación del sistema", level=2)
    add_paragraph(doc, "Ruta recomendada de instalación en Windows:")
    add_numbered(
        doc,
        [
            "Copiar el proyecto completo al equipo objetivo.",
            "Ejecutar scripts/install_windows.ps1 para crear el entorno virtual e instalar dependencias.",
            "Copiar .env.example a .env y definir credenciales de MySQL, APP_SECRET_KEY y contraseñas semilla.",
            "Ejecutar .\\.venv\\Scripts\\python.exe src\\main.py bootstrap para inicializar tablas y usuarios base.",
            "Iniciar la aplicación mediante scripts\\run_app.bat o python src/main.py.",
        ],
    )
    add_paragraph(doc, "Instalación manual alternativa (Linux/WSL):")
    add_numbered(
        doc,
        [
            "Crear entorno virtual con python -m venv .venv.",
            "Activar el entorno e instalar dependencias con pip install -r requirements.txt.",
            "Configurar el archivo .env con valores válidos.",
            "Ejecutar python src/main.py y validar el acceso web por el puerto 5000.",
        ],
    )
    doc.add_heading("4.2 Instalación de servicios", level=2)
    add_bullets(
        doc,
        [
            "MySQL: debe existir un esquema la_septima_estrella y un usuario con permisos sobre él.",
            "Docker: el proyecto incluye docker-compose.yml y Dockerfile para levantar aplicación y base de datos en conjunto.",
            "Playwright: se usa para smoke tests del frontend cuando se requiere validación de navegación real.",
        ],
    )
    add_paragraph(doc, "Errores comunes de instalación:")
    add_table(
        doc,
        ["Situación", "Causa probable", "Acción correctiva"],
        [
            ("La aplicación no inicia", "Variables .env incompletas o MySQL apagado.", "Verificar DB_HOST, credenciales y servicio MySQL."),
            ("No se puede iniciar sesión", "Usuarios semilla no creados.", "Ejecutar bootstrap o setup_users."),
            ("No cargan reportes", "Migraciones pendientes.", "Ejecutar python src/main.py migrate."),
            ("Falla entrenamiento IA", "No hay datos o falta modelo.", "Usar CSV demo o ejecutar setup_demo_ai."),
        ],
    )

    doc.add_heading("5. Requerimientos de hardware para operación", level=1)
    add_paragraph(
        doc,
        "Para operación continua se recomienda un equipo con procesador de cuatro núcleos, 16 GB de RAM, almacenamiento SSD y sistema operativo estable. En entornos con varios usuarios simultáneos conviene separar la base de datos del proceso web y calendarizar respaldos automáticos. El sistema puede crecer de forma controlada mediante contenedores, ajuste de la base de datos y eventual división del backend en servicios especializados.",
    )

    doc.add_heading("6. Bibliografía", level=1)
    add_bullets(
        doc,
        [
            "Pallets. (2026). Flask Documentation. https://flask.palletsprojects.com/",
            "PyMySQL Project. (2026). PyMySQL Documentation. https://pymysql.readthedocs.io/",
            "pandas Development Team. (2026). pandas documentation. https://pandas.pydata.org/docs/",
            "scikit-learn developers. (2026). scikit-learn documentation. https://scikit-learn.org/stable/",
            "Docker Inc. (2026). Docker Documentation. https://docs.docker.com/",
            "Microsoft. (2026). Playwright Documentation. https://playwright.dev/python/",
        ],
    )

    doc.save(TECH_DOC_PATH)


def build_user_manual() -> None:
    doc = Document()
    set_default_styles(doc)
    add_cover_page(doc, f"MANUAL DE USUARIO DEL SISTEMA {SYSTEM_NAME.upper()}")
    start_main_section(doc)
    add_toc(doc)
    doc.add_page_break()

    figure_counter = [0]

    doc.add_heading("1. Introducción", level=1)
    add_paragraph(
        doc,
        "El presente manual de usuario orienta al usuario final en el uso del sistema La Séptima Estrella. Su finalidad consiste en explicar de manera clara las opciones del menú, el uso de cada pantalla, la interpretación de mensajes y las acciones recomendadas ante fallos comunes. El documento no profundiza en detalles técnicos internos; se concentra en el manejo operativo del sistema por parte de perfiles administrativos y operativos.",
    )

    doc.add_heading("2. Descripción general del sistema", level=1)
    add_paragraph(
        doc,
        "La Séptima Estrella es un sistema web para gestión de ventas, control de productos, consulta de reportes, estadísticas con apoyo de inteligencia artificial, auditoría, administración de usuarios y respaldos del sistema. La solución atiende la necesidad de centralizar la operación de una tienda o negocio de barrio, manteniendo trazabilidad, control de inventario y soporte para decisiones de reposición.",
    )
    add_bullets(
        doc,
        [
            "Tipo de sistema: aplicación web de uso interno.",
            "Usuarios objetivo: administrador, gerente, vendedor y auditador.",
            "Procesos soportados: autenticación, ventas, inventario, reportes, estadísticas, respaldos y auditoría.",
        ],
    )

    doc.add_heading("3. Descripción de las opciones del menú", level=1)
    add_table(
        doc,
        ["Opción", "Ubicación funcional", "Propósito", "Restricción de acceso"],
        [
            ("Dashboard", "Menú lateral principal", "Presenta indicadores y alertas operativas.", ROLE_ACCESS["Dashboard"]),
            ("Registrar venta", "Menú lateral principal", "Gestiona el proceso de venta.", ROLE_ACCESS["Registrar venta"]),
            ("Productos", "Menú lateral principal", "Consulta y administra el catálogo.", ROLE_ACCESS["Productos"]),
            ("Reportes", "Menú lateral principal", "Consulta ventas por periodo y exporta resultados.", ROLE_ACCESS["Reportes"]),
            ("Estadísticas", "Menú lateral principal", "Analiza comportamiento histórico y pronóstico.", ROLE_ACCESS["Estadísticas"]),
            ("Auditoría", "Menú lateral principal", "Consulta eventos críticos del sistema.", ROLE_ACCESS["Auditoría"]),
            ("Usuarios", "Menú lateral principal", "Administra cuentas, roles y permisos.", ROLE_ACCESS["Usuarios"]),
            ("Respaldos", "Menú lateral principal", "Genera, descarga y restaura snapshots.", ROLE_ACCESS["Respaldos"]),
            ("Mi perfil", "Menú lateral principal", "Consulta capacidades y cierra sesión.", ROLE_ACCESS["Mi perfil"]),
        ],
    )

    doc.add_heading("4. Descripción y funcionalidad de las pantallas del sistema", level=1)
    add_paragraph(
        doc,
        "En este apartado se describen las pantallas principales del sistema, sus campos, acciones, flujo de uso y validaciones. La intención es que el usuario pueda ejecutar sus tareas sin apoyo externo.",
    )
    for section in SCREEN_SECTIONS:
        add_screen_section(doc, section, figure_counter)

    doc.add_heading("5. Descripción de los reportes del sistema", level=1)
    add_table(
        doc,
        ["Reporte", "Objetivo", "Información que contiene", "Filtros / salidas", "Uso para decisión"],
        [
            ("Reporte de ventas", "Consolidar ventas de un periodo.", "Ventas, total, usuario, método de pago y acceso a detalle.", "Rango de fechas; salida en pantalla, PDF y Excel.", "Permite sustentar ventas realizadas y comparar periodos."),
            ("Reporte estadístico", "Analizar comportamiento del catálogo.", "Productos más vendidos, cantidades, participación y tendencia.", "Rango de fechas; salida en pantalla, PDF y Excel.", "Ayuda a identificar productos de alta rotación."),
            ("Pronóstico de demanda", "Estimar cobertura y necesidad de compra.", "Promedio diario, riesgo, cobertura y compra sugerida.", "Rango de fechas, horizonte y filtros de riesgo.", "Apoya la reposición oportuna de inventario."),
            ("Auditoría", "Revisar trazabilidad del sistema.", "Fecha, evento, entidad, actor y detalle.", "Fecha, tipo, usuario; salida en pantalla, PDF y Excel.", "Permite control interno y verificación de acciones sensibles."),
            ("Respaldos", "Consultar evidencia de snapshots.", "Nombre de archivo, fecha, enlaces de descarga.", "Listado y descarga JSON/PDF.", "Ayuda a conservar evidencia y recuperar información."),
        ],
    )

    doc.add_heading("6. Interpretación de mensajes y errores", level=1)
    add_table(doc, ["Mensaje o situación", "Causa probable", "Acción recomendada"], COMMON_MESSAGES)

    doc.add_heading("7. Procedimiento a seguir en caso de fallos", level=1)
    add_numbered(
        doc,
        [
            "Leer completamente el mensaje mostrado por el sistema.",
            "Verificar si la sesión continúa activa y si el rango o formulario diligenciado es correcto.",
            "Intentar nuevamente solo cuando la información haya sido corregida.",
            "Si el problema persiste, registrar hora, módulo, mensaje observado y, de ser posible, una captura de pantalla.",
            "Contactar al administrador o responsable técnico con la evidencia recolectada.",
            "En caso de indisponibilidad prolongada, validar si el sistema está operando con respaldo local o si requiere restauración de snapshot.",
        ],
    )

    doc.add_heading("Glosario", level=1)
    add_table(
        doc,
        ["Término", "Descripción"],
        [
            ("Dashboard", "Pantalla inicial con indicadores operativos."),
            ("Snapshot", "Respaldo puntual del estado del sistema."),
            ("Cobertura", "Número estimado de días que alcanza el stock actual."),
            ("Compra sugerida", "Cantidad recomendada para reposición de inventario."),
            ("Deshabilitar venta", "Acción que revierte una venta sin borrar su trazabilidad."),
        ],
    )

    doc.add_heading("Anexos", level=1)
    add_paragraph(
        doc,
        "Las capturas incluidas en el manual corresponden al sistema real ejecutado localmente durante abril de 2026. Se recomienda actualizarlas si la interfaz cambia de forma significativa.",
    )

    doc.add_heading("Referencias", level=1)
    add_bullets(
        doc,
        [
            "Documentación interna del proyecto La Séptima Estrella.",
            "Instructivo para la elaboración del manual técnico 2026.",
            "Estructura instructiva del manual de usuario 2026.",
        ],
    )

    doc.save(USER_DOC_PATH)


def copy_outputs() -> None:
    shutil.copy2(TECH_DOC_PATH, TECH_DOWNLOAD_PATH)
    shutil.copy2(USER_DOC_PATH, USER_DOWNLOAD_PATH)


def main() -> None:
    ensure_dirs()
    diagrams = generate_diagrams()
    build_technical_manual(diagrams)
    build_user_manual()
    copy_outputs()
    manifest = {
        "technical_manual": str(TECH_DOC_PATH),
        "user_manual": str(USER_DOC_PATH),
        "technical_download_copy": str(TECH_DOWNLOAD_PATH),
        "user_download_copy": str(USER_DOWNLOAD_PATH),
    }
    (DOCS_DIR / "manuals_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
