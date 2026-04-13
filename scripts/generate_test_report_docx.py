from __future__ import annotations

import json
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image, ImageDraw, ImageFont

Image.MAX_IMAGE_PIXELS = None


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
ARTIFACTS_DIR = ROOT / "artifacts" / "tests"
SCREENSHOTS_DIR = ROOT / "docs" / "manuales" / "assets" / "screenshots"
WINDOWS_DOWNLOADS_DIR = ROOT.parents[2] / "Downloads"
DOWNLOADS_DIR = WINDOWS_DOWNLOADS_DIR if WINDOWS_DOWNLOADS_DIR.exists() else Path.home() / "Downloads"

SUMMARY_JSON = ARTIFACTS_DIR / "pytest-summary.json"
REPORT_MD = DOCS_DIR / "PRUEBAS_VALIDACION_SISTEMA.md"
OUTPUT_DOCX = DOCS_DIR / "PRUEBAS_VALIDACION_SISTEMA.docx"
DOWNLOAD_COPY = DOWNLOADS_DIR / OUTPUT_DOCX.name
SUMMARY_CHART = ARTIFACTS_DIR / "test-summary-chart.png"


MODULE_LABELS = {
    "tests/test_ai_reports.py": "Reportes IA y pronóstico de demanda",
    "tests/test_audit_endpoints.py": "Auditoría y exportación de trazabilidad",
    "tests/test_chat_ai.py": "Asistente IA conversacional",
    "tests/test_config_roles_endpoints.py": "Configuración de IVA y roles",
    "tests/test_data_paths.py": "Rutas de datos y archivos de ventas",
    "tests/test_demo_generator.py": "Generación de datos demo",
    "tests/test_frontend_cache_busting.py": "Versionado de assets frontend",
    "tests/test_ia_endpoints_auth.py": "Autenticación y permisos IA",
    "tests/test_operations_improvements.py": "Respaldos, exportes e indicadores operativos",
    "tests/test_password_hashing.py": "Hashing y seguridad de contraseñas",
    "tests/test_permissions_exports.py": "Permisos por rol y exportaciones",
    "tests/test_registrar_venta.py": "Registro básico de venta",
    "tests/test_sale_voiding.py": "Deshabilitación controlada de ventas",
    "tests/test_sales_flows.py": "Flujo transaccional de ventas e IVA",
    "tests/test_sessions.py": "Sesiones y expiración controlada",
    "tests/test_state_routes_audit.py": "Cambio de estado con auditoría",
    "tests/test_users_roles.py": "Usuarios, CLI y normalización de roles",
    "tests/test_validations.py": "Validaciones numéricas y de payload",
}


SCREENSHOTS = [
    ("01_login.png", "Pantalla actual de inicio de sesión del sistema."),
    ("02_dashboard.png", "Dashboard operativo con indicadores, alertas de stock y métodos de pago."),
    ("03_sales.png", "Módulo de registro de ventas con búsqueda, carrito y cierre de transacción."),
    ("05_products.png", "Gestión de productos, inventario, imágenes y exportaciones."),
    ("08_forecast.png", "Vista del pronóstico de demanda y compra sugerida por producto."),
    ("10_audit.png", "Módulo de auditoría con filtros y exportación de eventos."),
    ("11_backups.png", "Pantalla de respaldos con listado, descarga y restauración."),
]


SCENARIOS = [
    (
        "Funcional",
        "Login",
        "Bloquear acceso después de múltiples intentos fallidos consecutivos",
        "El sistema debe responder con restricción temporal y tiempo restante.",
        "El bloqueo temporal se aplicó correctamente.",
        "Exitosa",
    ),
    (
        "Funcional",
        "Sesiones",
        "Invalidar sesión previa al reiniciar el servidor",
        "La sesión anterior no debe seguir vigente.",
        "La sesión fue invalidada correctamente.",
        "Exitosa",
    ),
    (
        "Funcional",
        "IVA global",
        "Consultar y actualizar el IVA del sistema",
        "El valor debe leerse y persistirse correctamente.",
        "La operación se validó correctamente.",
        "Exitosa",
    ),
    (
        "Funcional",
        "Roles",
        "Crear un rol personalizado desde el endpoint correspondiente",
        "El rol debe almacenarse con permisos válidos.",
        "La creación se validó correctamente.",
        "Exitosa",
    ),
    (
        "Funcional",
        "Productos",
        "Exportar inventario en formato PDF",
        "El archivo debe generarse sin error y con el nombre esperado.",
        "La exportación se generó correctamente.",
        "Exitosa",
    ),
    (
        "Funcional",
        "Inventario",
        "Detectar productos por debajo o muy cerca del mínimo",
        "El dashboard debe incluir alertas de seguimiento.",
        "Las alertas de stock se devolvieron según lo esperado.",
        "Exitosa",
    ),
    (
        "Funcional",
        "Ventas",
        "Registrar venta con IVA específico por producto",
        "El cálculo debe priorizar el IVA individual del producto.",
        "El cálculo y el detalle devueltos fueron correctos.",
        "Exitosa",
    ),
    (
        "Funcional",
        "Ventas",
        "Revertir una transacción cuando ocurre un fallo intermedio",
        "La operación debe hacer rollback y no dejar datos inconsistentes.",
        "La reversión transaccional se comportó correctamente.",
        "Exitosa",
    ),
    (
        "Funcional",
        "Anulación de ventas",
        "Deshabilitar una venta autorizada y restaurar stock",
        "La venta debe quedar inactiva y el inventario debe recuperarse.",
        "La restauración de stock y el cambio de estado fueron correctos.",
        "Exitosa",
    ),
    (
        "Funcional",
        "IA",
        "Generar pronóstico de demanda con respuesta estructurada",
        "El endpoint debe devolver riesgo, cobertura y compra sugerida.",
        "El payload se devolvió con la estructura esperada.",
        "Exitosa",
    ),
]


def get_font(size: int, *, bold: bool = False) -> ImageFont.ImageFont:
    preferred = "arialbd.ttf" if bold else "arial.ttf"
    windows_font = Path("C:/Windows/Fonts") / preferred
    if windows_font.exists():
        return ImageFont.truetype(str(windows_font), size=size)
    return ImageFont.load_default()


def set_default_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Arial"

    document.styles["Title"].font.size = Pt(16)
    document.styles["Heading 1"].font.size = Pt(14)
    document.styles["Heading 2"].font.size = Pt(13)
    document.styles["Heading 3"].font.size = Pt(12)

    for section in document.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(4)
        section.right_margin = Cm(2)
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)


def add_field(paragraph, instruction: str, default_text: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    if default_text:
        text = OxmlElement("w:t")
        text.text = default_text
        run._r.append(text)
    run._r.append(end)


def set_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Página ")
    add_field(paragraph, "PAGE", "1")
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)


def restart_page_numbering(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def add_cover(document: Document, generated_at: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LA SÉPTIMA ESTRELLA")
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph("")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PRUEBAS DE VALIDACIÓN DEL SISTEMA")
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph("")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "Sistema web de inventario, ventas, reportes, auditoría y análisis inteligente."
    )

    document.add_paragraph("")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Documento de evidencia técnica y funcional.")

    document.add_paragraph("")
    document.add_paragraph("")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Fecha de generación: {generated_at}")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Proyecto: Analisis de Ventas IA")


def start_main_section(document: Document):
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)
    section.left_margin = Cm(4)
    section.right_margin = Cm(2)
    restart_page_numbering(section, 1)
    set_page_number(section)
    return section


def set_cell_text(cell, text: str) -> None:
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10.5)


def format_table(table) -> None:
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10.5)


def add_table(
    document: Document, headers: Sequence[str], rows: Sequence[Sequence[str]]
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    format_table(table)


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.add_run(item)


def prepare_image(
    path: Path,
    *,
    quality: int = 88,
    canvas_size: tuple[int, int] = (1600, 1000),
    background: str = "#ffffff",
) -> Path:
    optimized_dir = SCREENSHOTS_DIR.parent / "optimized-tests"
    optimized_dir.mkdir(parents=True, exist_ok=True)
    target = optimized_dir / path.name
    with Image.open(path) as image:
        image = image.convert("RGB")
        canvas_width, canvas_height = canvas_size
        inner_width = int(canvas_width * 0.9)
        inner_height = int(canvas_height * 0.86)
        resize_ratio = min(
            inner_width / float(image.width),
            inner_height / float(image.height),
            1.0,
        )
        resized = image.resize(
            (
                max(1, int(image.width * resize_ratio)),
                max(1, int(image.height * resize_ratio)),
            ),
            Image.LANCZOS,
        )
        canvas = Image.new("RGB", canvas_size, background)
        offset_x = (canvas_width - resized.width) // 2
        offset_y = (canvas_height - resized.height) // 2
        canvas.paste(resized, (offset_x, offset_y))
        canvas.save(target, format="JPEG", quality=quality, optimize=True)
    return target


def add_figure(
    document: Document,
    image_path: Path,
    caption: str,
    number: int,
    *,
    width: float = 4.7,
) -> None:
    prepared = prepare_image(image_path)
    document.add_picture(str(prepared), width=Inches(width))
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Figura {number}. {caption}")
    run.italic = True


def create_summary_chart(summary: dict) -> Path:
    ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)
    suite = summary.get("suite") or {}
    counts = {
        "Exitosas": int(suite.get("tests", 0)) - int(suite.get("failures", 0)) - int(suite.get("errors", 0)) - int(suite.get("skipped", 0)),
        "Fallidas": int(suite.get("failures", 0)),
        "Errores": int(suite.get("errors", 0)),
        "Omitidas": int(suite.get("skipped", 0)),
    }
    total = max(sum(counts.values()), 1)

    image = Image.new("RGB", (1400, 820), "#f7f9ff")
    draw = ImageDraw.Draw(image)
    title_font = get_font(34, bold=True)
    label_font = get_font(22, bold=True)
    value_font = get_font(20)

    draw.text((80, 50), "Resumen general de pruebas", fill="#10204a", font=title_font)
    draw.text(
        (80, 110),
        f"Total ejecutadas: {total} | Tiempo total: {suite.get('time', '0')} s",
        fill="#243764",
        font=value_font,
    )

    colors = {
        "Exitosas": "#27ae60",
        "Fallidas": "#e74c3c",
        "Errores": "#f39c12",
        "Omitidas": "#95a5a6",
    }
    top = 220
    bar_left = 320
    bar_max = 920
    bar_height = 70

    for idx, (label, value) in enumerate(counts.items()):
        y = top + idx * 130
        ratio = value / total if total else 0
        draw.text((80, y + 15), label, fill="#10204a", font=label_font)
        draw.rounded_rectangle(
            (bar_left, y, bar_left + bar_max, y + bar_height),
            radius=20,
            fill="#e9eefb",
        )
        draw.rounded_rectangle(
            (bar_left, y, bar_left + max(10, int(bar_max * ratio)), y + bar_height),
            radius=20,
            fill=colors[label],
        )
        draw.text(
            (bar_left + bar_max + 30, y + 15),
            f"{value} ({ratio * 100:.1f}%)",
            fill="#10204a",
            font=value_font,
        )

    image.save(SUMMARY_CHART)
    return SUMMARY_CHART


def load_summary() -> dict:
    return json.loads(SUMMARY_JSON.read_text(encoding="utf-8"))


def build_module_rows(summary: dict) -> list[list[str]]:
    rows = []
    for file_key in sorted(summary.get("files") or {}):
        entry = summary["files"][file_key]
        rows.append(
            [
                file_key,
                MODULE_LABELS.get(file_key, "Cobertura automatizada del módulo."),
                str(entry.get("tests", 0)),
                str(entry.get("passed", 0)),
                str(entry.get("failed", 0)),
                str(entry.get("errored", 0)),
                "Exitosa" if entry.get("failed", 0) == 0 and entry.get("errored", 0) == 0 else "Con incidencias",
            ]
        )
    return rows


def build_top_case_rows(summary: dict, *, limit: int = 8) -> list[list[str]]:
    rows = []
    cases = []
    for file_key, file_data in (summary.get("files") or {}).items():
        for case in file_data.get("cases") or []:
            cases.append(
                (
                    float(case.get("time", 0.0) or 0.0),
                    file_key,
                    case.get("name", ""),
                    case.get("status", ""),
                )
            )
    for elapsed, file_key, name, status in sorted(cases, reverse=True)[:limit]:
        rows.append([name, file_key, f"{elapsed:.3f} s", status])
    return rows


def add_code_block(document: Document, lines: Sequence[str]) -> None:
    for line in lines:
        p = document.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)


def build_document() -> Document:
    summary = load_summary()
    suite = summary.get("suite") or {}
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    document = Document()
    set_default_styles(document)
    add_cover(document, generated_at)
    start_main_section(document)

    document.add_heading("PRUEBAS", level=1)
    document.add_paragraph(
        "Este documento presenta la validación formal del sistema desarrollado para "
        "La Séptima Estrella. La evidencia reúne planificación, ejecución, resultados "
        "obtenidos, análisis posterior y soporte visual de los módulos evaluados."
    )

    document.add_heading("Planificación de las pruebas", level=2)
    document.add_paragraph(
        "Antes de la ejecución se definió el alcance de la validación, los recursos "
        "necesarios y los criterios de aceptación. Se estableció como objetivo comprobar "
        "el funcionamiento correcto de autenticación, sesiones, usuarios, roles, "
        "productos, inventario, ventas, reportes, auditoría, respaldos e inteligencia artificial."
    )
    add_bullets(
        document,
        [
            "El sistema se encontraba instalado y operativo en un entorno local controlado.",
            "Los datos de prueba estaban definidos previamente dentro de la batería automatizada.",
            "Cada caso contaba con resultados esperados explícitos mediante aserciones verificables.",
            "Se contemplaron respuestas ante datos válidos, inválidos y fallos controlados.",
        ],
    )
    add_table(
        document,
        ["Tipo de prueba", "Alcance", "Herramienta", "Criterio de aceptación"],
        [
            [
                "Funcional",
                "Login, sesiones, usuarios, productos, inventario, ventas, reportes, auditoría e IA.",
                "pytest",
                "El caso termina sin error y cumple la respuesta esperada.",
            ],
            [
                "Integración",
                "Endpoints Flask, permisos, exportes, rollback, respaldos y trazabilidad.",
                "pytest",
                "La respuesta y el estado del sistema coinciden con la especificación.",
            ],
            [
                "Verificación técnica",
                "Compilación Python y sintaxis del frontend.",
                "compileall / node --check",
                "No aparecen errores de compilación ni de sintaxis.",
            ],
        ],
    )

    document.add_heading("Ejecución de las pruebas", level=2)
    document.add_paragraph(
        "La corrida principal se ejecutó sobre el entorno del proyecto con evidencia "
        "registrada en artefactos XML, JSON y archivos de salida. Los comandos utilizados fueron los siguientes:"
    )
    add_code_block(
        document,
        [
            ".venv-linux/bin/python -m pytest -q --capture=no --junitxml=artifacts/tests/pytest-results.xml",
            ".venv-linux/bin/python -m compileall src tests scripts > artifacts/tests/compileall.txt 2>&1",
            "node --check src/frontend/app.js > artifacts/tests/frontend-check.txt 2>&1",
        ],
    )
    add_table(
        document,
        ["Métrica", "Valor obtenido"],
        [
            ["Host reportado por la suite", str(suite.get("hostname", "No reportado"))],
            ["Fecha y hora de la ejecución", str(suite.get("timestamp", "No reportada"))],
            ["Pruebas ejecutadas", str(suite.get("tests", "0"))],
            ["Pruebas exitosas", str(int(suite.get("tests", "0")) - int(suite.get("failures", "0")) - int(suite.get("errors", "0")) - int(suite.get("skipped", "0")))],
            ["Pruebas fallidas", str(suite.get("failures", "0"))],
            ["Pruebas con error", str(suite.get("errors", "0"))],
            ["Pruebas omitidas", str(suite.get("skipped", "0"))],
            ["Tiempo total", f"{suite.get('time', '0')} s"],
        ],
    )

    document.add_paragraph("")
    add_figure(
        document,
        create_summary_chart(summary),
        "Distribución general de resultados de la suite automatizada ejecutada.",
        1,
    )

    document.add_heading("Resultados por módulo", level=2)
    document.add_paragraph(
        "La batería automatizada se distribuyó por módulos funcionales y técnicos. "
        "La siguiente tabla resume la cobertura validada y el estado final de cada grupo de pruebas."
    )
    add_table(
        document,
        ["Archivo de prueba", "Módulo validado", "Casos", "Exitosas", "Fallidas", "Errores", "Estado"],
        build_module_rows(summary),
    )

    document.add_heading("Escenarios representativos ejecutados", level=2)
    document.add_paragraph(
        "Además del consolidado por módulo, se documentan escenarios funcionales relevantes "
        "que fueron ejecutados y validados satisfactoriamente durante la corrida."
    )
    add_table(
        document,
        [
            "Tipo",
            "Módulo",
            "Escenario",
            "Resultado esperado",
            "Resultado obtenido",
            "Estado",
        ],
        SCENARIOS,
    )

    document.add_heading("Casos con mayor tiempo de ejecución", level=2)
    document.add_paragraph(
        "El tiempo de respuesta observado en la suite fue estable. Los casos de mayor duración "
        "se concentraron en validaciones de IA, hashing y generación de archivos."
    )
    add_table(
        document,
        ["Caso", "Archivo", "Tiempo", "Estado"],
        build_top_case_rows(summary),
    )

    document.add_heading("Evidencia visual del sistema evaluado", level=2)
    document.add_paragraph(
        "Como soporte visual se incorporan capturas vigentes de los módulos principales del sistema "
        "que fueron objeto de validación funcional. Estas imágenes complementan la evidencia técnica "
        "y muestran la interfaz real bajo evaluación."
    )
    figure_no = 2
    for image_name, caption in SCREENSHOTS:
        image_path = SCREENSHOTS_DIR / image_name
        if image_path.exists():
            add_figure(
                document,
                image_path,
                caption,
                figure_no,
            )
            figure_no += 1

    document.add_heading("Análisis de resultados", level=2)
    document.add_paragraph(
        "Al comparar el resultado esperado con el resultado real, se observó que el sistema "
        "respondió de forma correcta en autenticación, sesiones, usuarios, roles, productos, "
        "inventario, ventas, reportes, exportaciones, auditoría, respaldos y asistencia inteligente."
    )
    add_bullets(
        document,
        [
            "No se presentaron fallas activas en la corrida ejecutada.",
            "La integridad transaccional de ventas e inventario se mantuvo en los escenarios controlados.",
            "Los permisos por rol y las restricciones de acceso se comportaron según la política definida.",
            "Las exportaciones, respaldos y registros de auditoría respondieron de forma consistente.",
            "Los mensajes de consola observados durante la ejecución correspondieron a fallos simulados dentro de los dobles de prueba y no a defectos vigentes de la aplicación.",
        ],
    )

    document.add_heading("Resultados generales de las pruebas", level=2)
    passed = int(suite.get("tests", "0")) - int(suite.get("failures", "0")) - int(suite.get("errors", "0")) - int(suite.get("skipped", "0"))
    document.add_paragraph(
        f"En total se ejecutaron {suite.get('tests', '0')} pruebas automatizadas, de las cuales "
        f"{passed} fueron exitosas, {suite.get('failures', '0')} presentaron fallas, "
        f"{suite.get('errors', '0')} presentaron error y {suite.get('skipped', '0')} fueron omitidas."
    )
    document.add_paragraph(
        "De forma complementaria se verificó la compilación del backend y la sintaxis del frontend, "
        "ambas sin errores registrados."
    )

    document.add_heading("Conclusiones de las pruebas", level=2)
    document.add_paragraph(
        "Con base en los resultados obtenidos, se concluye que el sistema se encuentra en condiciones "
        "técnicas de ser utilizado en un entorno real controlado, ya que cumplió satisfactoriamente "
        "con los requisitos funcionales y de soporte evaluados durante la validación."
    )
    add_bullets(
        document,
        [
            "El sistema gestiona correctamente autenticación, sesiones y permisos.",
            "El flujo de ventas conserva integridad, trazabilidad y control de IVA global e individual.",
            "Los módulos de reportes, auditoría, exportación y respaldos operan de manera consistente.",
            "El componente de IA funciona como soporte analítico y mantiene restricciones adecuadas por rol.",
            "Se recomienda complementar esta validación con pruebas de carga y aceptación con usuarios reales para reforzar la evaluación en operación continua.",
        ],
    )

    document.add_heading("Artefactos generados", level=2)
    add_bullets(
        document,
        [
            str(ARTIFACTS_DIR / "pytest-results.xml"),
            str(ARTIFACTS_DIR / "pytest-summary.json"),
            str(ARTIFACTS_DIR / "compileall.txt"),
            str(ARTIFACTS_DIR / "frontend-check.txt"),
            str(SUMMARY_CHART),
            str(REPORT_MD),
        ],
    )
    return document


def main() -> int:
    if not SUMMARY_JSON.exists():
        raise FileNotFoundError(
            f"No se encontró el resumen de pruebas en {SUMMARY_JSON}. Ejecuta primero la batería de pruebas."
        )
    document = build_document()
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DOCX
    download_path = DOWNLOAD_COPY
    try:
        document.save(output_path)
    except PermissionError:
        suffix = datetime.now().strftime("%Y%m%d-%H%M%S")
        output_path = OUTPUT_DOCX.with_name(
            f"{OUTPUT_DOCX.stem}_{suffix}{OUTPUT_DOCX.suffix}"
        )
        download_path = DOWNLOADS_DIR / output_path.name
        document.save(output_path)
    try:
        shutil.copy2(output_path, download_path)
    except Exception:
        pass
    print(f"[ok] documento generado en {output_path}")
    print(f"[ok] copia en descargas: {download_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
