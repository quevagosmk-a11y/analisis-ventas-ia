import os
import sys
from datetime import datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

ROOT_DIR = os.path.dirname(os.path.dirname(__file__))
SRC_DIR = os.path.join(ROOT_DIR, "src")
if SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

from src import main as app_module


class DummyConnection:
    def close(self):
        pass


def _build_client(monkeypatch, role: str):
    monkeypatch.setattr(
        app_module,
        "get_session_state",
        lambda *args, **kwargs: app_module.SessionState(True, False, False, None),
    )
    app_module.app.config["TESTING"] = True
    client = app_module.app.test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["role"] = role
        sess["username"] = role
        sess["name"] = role.title()
        sess["session_token"] = "testing-token"
        sess["idle_minutes"] = 45
        sess["last_activity"] = datetime.now(timezone.utc).isoformat()
    return client


@pytest.fixture()
def auditor_client(monkeypatch):
    with _build_client(monkeypatch, "auditador") as client:
        yield client


@pytest.fixture()
def seller_client(monkeypatch):
    with _build_client(monkeypatch, "vendedor") as client:
        yield client


def test_audit_endpoint_requires_audit_permission(monkeypatch, seller_client):
    monkeypatch.setattr(
        app_module,
        "get_db_conn",
        lambda: (_ for _ in ()).throw(AssertionError("No debe consultar BD")),
    )
    response = seller_client.get("/api/auditoria")
    assert response.status_code == 403
    payload = response.get_json()
    assert "error" in (payload or {})


def test_audit_endpoint_merges_current_and_legacy_entries(monkeypatch, auditor_client):
    monkeypatch.setattr(app_module, "get_db_conn", lambda: DummyConnection())
    monkeypatch.setattr(
        app_module,
        "fetch_audit_event_rows",
        lambda _conn, *, limit: [
            {
                "id": 101,
                "event_type": "producto_actualizado",
                "entity_type": "producto",
                "entity_id": 8,
                "description": "Producto actualizado",
                "details": {},
                "actor_user_id": 1,
                "actor_username": "admin",
                "actor_name": "Administrador",
                "created_at": "2026-03-19T10:00:00-05:00",
            }
        ],
    )
    monkeypatch.setattr(
        app_module,
        "fetch_legacy_price_audit_rows",
        lambda _conn, *, limit: [
            {
                "id": 55,
                "event_type": "precio_actualizado",
                "entity_type": "producto",
                "entity_id": 8,
                "description": "Precio actualizado en Cafe",
                "details": {
                    "precio_anterior": 10000.0,
                    "precio_nuevo": 12000.0,
                },
                "actor_user_id": 1,
                "actor_username": "admin",
                "actor_name": "Administrador",
                "created_at": "2026-03-18T09:00:00-05:00",
            }
        ],
    )
    monkeypatch.setattr(
        app_module,
        "fetch_legacy_inventory_audit_rows",
        lambda _conn, *, limit: [
            {
                "id": 12,
                "event_type": "stock_ajustado",
                "entity_type": "producto",
                "entity_id": 8,
                "description": "Inventario ajuste en Cafe",
                "details": {
                    "cantidad": -2,
                    "stock_anterior": 10,
                    "stock_nuevo": 8,
                },
                "actor_user_id": 2,
                "actor_username": "bodega",
                "actor_name": "Bodega",
                "created_at": "2026-03-17T08:00:00-05:00",
            }
        ],
    )

    response = auditor_client.get("/api/auditoria?limit=10")
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["success"] is True
    assert payload["summary"]["total"] == 3
    assert [row["event_type"] for row in payload["eventos"]] == [
        "producto_actualizado",
        "precio_actualizado",
        "stock_ajustado",
    ]


def test_audit_endpoint_filters_alias_event_names(monkeypatch, auditor_client):
    monkeypatch.setattr(app_module, "get_db_conn", lambda: DummyConnection())
    monkeypatch.setattr(
        app_module,
        "fetch_audit_event_rows",
        lambda _conn, *, limit: [
            {
                "id": 7,
                "event_type": "producto_minimo_actualizado",
                "entity_type": "producto",
                "entity_id": 4,
                "description": "Stock minimo actualizado",
                "details": {},
                "actor_user_id": 1,
                "actor_username": "admin",
                "actor_name": "Administrador",
                "created_at": "2026-03-19T11:00:00-05:00",
            }
        ],
    )
    monkeypatch.setattr(
        app_module, "fetch_legacy_price_audit_rows", lambda _conn, *, limit: []
    )
    monkeypatch.setattr(
        app_module, "fetch_legacy_inventory_audit_rows", lambda _conn, *, limit: []
    )

    response = auditor_client.get(
        "/api/auditoria?action=minimo_stock_actualizado&user=admin"
    )
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["summary"]["total"] == 1
    assert payload["eventos"][0]["entity_id"] == 4


def test_audit_endpoint_hides_export_events(monkeypatch, auditor_client):
    monkeypatch.setattr(app_module, "get_db_conn", lambda: DummyConnection())
    monkeypatch.setattr(
        app_module,
        "fetch_audit_event_rows",
        lambda _conn, *, limit: [
            {
                "id": 9,
                "event_type": "auditoria_exportada",
                "entity_type": "reporte",
                "entity_id": None,
                "description": "Exportación de auditoría",
                "details": {},
                "actor_user_id": 1,
                "actor_username": "admin",
                "actor_name": "Administrador",
                "created_at": "2026-03-19T12:00:00-05:00",
            },
            {
                "id": 8,
                "event_type": "producto_actualizado",
                "entity_type": "producto",
                "entity_id": 2,
                "description": "Producto actualizado",
                "details": {},
                "actor_user_id": 1,
                "actor_username": "admin",
                "actor_name": "Administrador",
                "created_at": "2026-03-19T11:00:00-05:00",
            },
        ],
    )
    monkeypatch.setattr(
        app_module, "fetch_legacy_price_audit_rows", lambda _conn, *, limit: []
    )
    monkeypatch.setattr(
        app_module, "fetch_legacy_inventory_audit_rows", lambda _conn, *, limit: []
    )

    response = auditor_client.get("/api/auditoria?limit=10")
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["summary"]["total"] == 1
    assert [row["event_type"] for row in payload["eventos"]] == [
        "producto_actualizado"
    ]


def test_audit_export_endpoint_returns_excel(monkeypatch, auditor_client):
    monkeypatch.setattr(
        app_module,
        "load_audit_trail_payload",
        lambda **_kwargs: {
            "success": True,
            "eventos": [
                {
                    "id": 9,
                    "created_at": "2026-03-18T10:30:00-05:00",
                    "event_type": "precio_actualizado",
                    "description": "Precio actualizado en Cafe",
                    "entity_type": "producto",
                    "entity_id": 3,
                    "actor_user_id": 1,
                    "actor_username": "admin",
                    "actor_name": "Administrador",
                }
            ],
            "summary": {
                "total": 1,
                "periodo": {"desde": "2026-03-01", "hasta": "2026-03-19"},
            },
        },
    )

    response = auditor_client.get(
        "/api/auditoria/export?from=2026-03-01&to=2026-03-19&format=excel"
    )
    assert response.status_code == 200
    assert (
        response.mimetype
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert len(response.data) > 20
    disposition = response.headers.get("Content-Disposition") or ""
    assert "reporte-auditoria-2026-03-01-a-2026-03-19.xlsx" in disposition


def test_build_audit_export_file_uses_main_label_helpers(monkeypatch):
    monkeypatch.setattr(app_module, "build_audit_entity_label", lambda _entry: "Entidad custom")
    monkeypatch.setattr(app_module, "build_audit_actor_label", lambda _entry: "Usuario custom")

    output, mimetype, extension = app_module.build_audit_export_file(
        {
            "eventos": [
                {
                    "id": 9,
                    "created_at": "2026-03-18T10:30:00-05:00",
                    "event_type": "precio_actualizado",
                    "description": "Precio actualizado en Cafe",
                    "entity_type": "producto",
                    "entity_id": 3,
                    "actor_user_id": 1,
                    "actor_username": "admin",
                    "actor_name": "Administrador",
                }
            ],
            "summary": {"total": 1},
        },
        "excel",
    )

    workbook = load_workbook(filename=BytesIO(output.getvalue()))
    rows = list(workbook.active.iter_rows(values_only=True))
    assert mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert extension == "xlsx"
    assert ("Entidad custom", "Usuario custom") in [
        (row[4], row[5]) for row in rows if len(row) >= 6
    ]
